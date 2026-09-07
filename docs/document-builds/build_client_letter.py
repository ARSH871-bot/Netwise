"""Client confirmation letter for Senaka to sign (Group Report Appendix C).

Written so it can honestly be signed, or honestly be amended. The factual
paragraphs are things he can check. The satisfaction paragraph is his own
judgement and is marked as editable, with a strike-out instruction, because a
letter somebody felt obliged to sign is worth nothing to a marker and is not
fair to him.

The limitations paragraph is deliberately included. Asking a client to endorse
a product without mentioning what it cannot do would be asking him to sign
something we ourselves do not claim.
"""
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

DOCS = pathlib.Path(__file__).resolve().parents[1]
OUT = DOCS / "Client Confirmation Letter - for Senaka to sign.docx"

d = Document()
st = d.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.2
for name, size in (("Heading 1", 15), ("Heading 2", 12)):
    s = d.styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string("1F3B57")
    s.font.bold = True


def p(t, size=None, italic=False, grey=False, bold=False):
    par = d.add_paragraph()
    r = par.add_run(t)
    if size:
        r.font.size = Pt(size)
    r.font.italic = italic
    r.font.bold = bold
    if grey:
        r.font.color.rgb = RGBColor.from_string("5A6472")
    return par


def note(t):
    par = d.add_paragraph()
    r = par.add_run(t)
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor.from_string("B4231F")
    par.paragraph_format.left_indent = Inches(0.2)
    par.paragraph_format.space_after = Pt(12)


# --------------------------------------------------------------- note to Arsh
note("NOTE FOR ARSH \u2014 delete this box before sending. Send this to Senaka as "
     "an editable file, not a PDF. He should feel free to change any wording, "
     "strike out anything he does not agree with, or write his own letter "
     "instead. A letter he has edited is better evidence than one he has only "
     "signed. If he would rather put it on university letterhead, that is "
     "better still.")

# --------------------------------------------------------------- letter
p("Client Confirmation \u2014 Netwise", size=15, bold=True)
p("IA728001 Studio 5, Block 3, 2026 \u2014 Auckland International Campus",
  size=10.5, grey=True)
d.add_paragraph()

p("To whom it may concern,")

p("I acted as the client for the Netwise project during Studio 5, Block 3, "
  "2026. The team was Arsh Vhora, Ankeet Patel, Shubham Kataria and Samika "
  "Perera.")

p("Netwise is a tool that reads exported network device configuration files, "
  "finds security misconfigurations in them, and explains those findings in "
  "plain English. It runs entirely offline and does not connect to a live "
  "network.")

p("During the project I:", bold=True)
for item in [
    "met the team to give the original requirements, including the two "
    "directions of translation I asked for \u2014 turning technical analysis "
    "output into plain English, and turning a plain-English instruction into "
    "a proposed configuration change;",
    "provided an anonymised export of my own pfSense firewall configuration "
    "for the team to test against;",
    "set the constraint that the tool must never push a change to a live "
    "device, and that configuration data must not be sent to any external "
    "service;",
    "gave feedback and correction at several points during the block, "
    "including on the scope of what was realistic.",
]:
    par = d.add_paragraph(item, style="List Bullet")
    par.paragraph_format.space_after = Pt(4)

d.add_paragraph()
p("I have seen the tool demonstrated. The team analysed a configuration file, "
  "showed the findings with the evidence behind each one, showed the "
  "plain-English explanations, and exported the result as a report.")

p("I am satisfied with what was delivered against what I asked for.", bold=True)
note("SENAKA \u2014 the sentence above is your judgement, not the team's. Please "
     "change it, qualify it, or delete it if it does not reflect your view. "
     "The rest of this letter is a record of fact and the team would rather "
     "have an accurate letter than a positive one.")

p("The team were clear with me about what the tool does not yet do. In "
  "particular, two of its three analyses do not yet read a security policy "
  "written by the user, and the converter for my own firewall format cannot "
  "yet handle every construct in my configuration \u2014 it stops and says which "
  "ones instead of guessing. Four questions about my firewall are still with "
  "me for a decision. I regard both of these as honestly reported rather than "
  "hidden.")

d.add_paragraph()
p("Signed,")
d.add_paragraph()

t = d.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
for label, value in [
    ("Name", "Senaka Amarakeerthi"),
    ("Position", "Senior Lecturer"),
    ("Signature", "________________________________"),
    ("Date", "________________________________"),
]:
    cells = t.add_row().cells
    for i, txt in enumerate((label, value)):
        cells[i].text = ""
        run = cells[i].paragraphs[0].add_run(txt)
        run.font.size = Pt(11)
        if i == 0:
            run.bold = True
    cells[0].width = Inches(1.3)
    cells[1].width = Inches(3.6)

OUT.parent.mkdir(parents=True, exist_ok=True)
# `python -P` drops the script's own directory from sys.path -- that is
# exactly what it is for -- so a sibling import has to be explicit.
import sys as _sys, pathlib as _pl                          # noqa: E402
_sys.path.insert(0, str(_pl.Path(__file__).resolve().parent))
from _guard import refuse_if_edited                         # noqa: E402
import _guard                                               # noqa: E402
refuse_if_edited(OUT, _guard.blocks(d))
d.save(str(OUT))
print(f"wrote: {OUT.name}")
print(f"  paragraphs {len(d.paragraphs)}  "
      f"~{sum(len(x.text.split()) for x in d.paragraphs)} words")
