"""Check the deck fits: no shape off-slide, no text taller than its box.

Cannot render (no LibreOffice), so measure instead. Line wrapping is estimated
with PIL's real Calibri metrics where available, which is far closer than a
characters-per-inch guess.
"""
import pathlib
from pptx import Presentation
from pptx.util import Emu

try:
    from PIL import ImageFont
    F = r"C:\Windows\Fonts\calibri.ttf"
    FB = r"C:\Windows\Fonts\calibrib.ttf"
    _cache = {}
    def width_pt(txt, size, bold):
        key = (round(size), bold)
        if key not in _cache:
            _cache[key] = ImageFont.truetype(FB if bold else F, int(size * 4))
        return _cache[key].getlength(txt) / 4.0
    MODE = "PIL/Calibri"
except Exception as e:                                    # noqa: BLE001
    def width_pt(txt, size, bold):
        return len(txt) * size * (0.52 if bold else 0.48)
    MODE = f"estimate ({e})"

SW, SH = 13.333, 7.5
DOCS = pathlib.Path(__file__).resolve().parents[1]
prs = Presentation(str(DOCS / "Demo" / "Netwise - Class Presentation.pptx"))
print(f"  wrapping measured with: {MODE}\n")

problems = []
for n, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        L, T = Emu(sh.left).inches, Emu(sh.top).inches
        W, H = Emu(sh.width).inches, Emu(sh.height).inches
        if L < -0.02 or T < -0.02 or L + W > SW + 0.02 or T + H > SH + 0.02:
            problems.append(f"slide {n}: {sh.shape_type} off-slide "
                            f"({L:.2f},{T:.2f}) {W:.2f}x{H:.2f}")
        if not sh.has_text_frame:
            continue
        # 0.1" default internal margin each side
        avail_pt = (W - 0.2) * 72
        total_pt = 0.0
        for para in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in para.runs)
            if not txt:
                total_pt += 12
                continue
            size = max((r.font.size.pt for r in para.runs if r.font.size), default=18)
            bold = any(r.font.bold for r in para.runs)
            ls = para.line_spacing or 1.0
            w = width_pt(txt, size, bold)
            nlines = max(1, -(-int(w * 100) // int(avail_pt * 100)))
            total_pt += nlines * size * 1.21 * ls
            total_pt += (para.space_after.pt if para.space_after else 0)
        need = total_pt / 72 + 0.1
        if need > H + 0.03:
            head = "".join(r.text for r in sh.text_frame.paragraphs[0].runs)[:44]
            problems.append(f"slide {n}: text needs {need:.2f}\" in a {H:.2f}\" box"
                            f"  -- \"{head}\"")

notes = sum(1 for s in prs.slides if s.has_notes_slide
            and s.notes_slide.notes_text_frame.text.strip())
print(f"  {len(prs.slides)} slides, {notes} with speaker notes")
if problems:
    print(f"\n  {len(problems)} LAYOUT PROBLEM(S):")
    for p in problems:
        print(f"    - {p}")
else:
    print("  no overflow, nothing off-slide")
